"""Excel persistence layer. All data access goes through here."""

import os  
import re  
import shutil  
import threading  
from datetime import datetime

import openpyxl

import config

_lock = threading.RLock()  
_wb = None  
_mtime = None  
_exercise_cache = None

_SHEETS = (  
    (config.SHEET_LOG, config.LOG_HEADERS),  
    (config.SHEET_PRS, config.PR_HEADERS),  
    (config.SHEET_PROGRAM, config.PROGRAM_HEADERS),  
    (config.SHEET_ALIASES, config.ALIAS_HEADERS),  
)

BUILTIN_ALIASES = {  
    "bench": "Bench Press", "bench press": "Bench Press", "flat bench": "Bench Press",  
    "incline": "Incline Bench Press", "incline bench": "Incline Bench Press",  
    "incline press": "Incline Bench Press",  
    "squat": "Back Squat", "back squat": "Back Squat", "front squat": "Front Squat",  
    "deadlift": "Deadlift", "dead lift": "Deadlift",  
    "rdl": "Romanian Deadlift", "romanian deadlift": "Romanian Deadlift",  
    "sumo": "Sumo Deadlift", "sumo deadlift": "Sumo Deadlift",  
    "ohp": "Overhead Press", "overhead press": "Overhead Press",  
    "military press": "Overhead Press", "shoulder press": "Overhead Press",  
    "pull up": "Pull Up", "pullup": "Pull Up",  
    "chin up": "Chin Up", "chinup": "Chin Up",  
    "row": "Barbell Row", "barbell row": "Barbell Row", "bent over row": "Barbell Row",  
    "dumbbell row": "Dumbbell Row", "db row": "Dumbbell Row",  
    "curl": "Barbell Curl", "bicep curl": "Barbell Curl",  
    "dip": "Dip", "leg press": "Leg Press", "lunge": "Lunge",  
    "hip thrust": "Hip Thrust", "calf raise": "Calf Raise",  
    "lat pulldown": "Lat Pulldown", "pulldown": "Lat Pulldown",  
    "tricep extension": "Tricep Extension", "skullcrusher": "Skullcrusher",  
    "leg curl": "Leg Curl", "leg extension": "Leg Extension",  
    "face pull": "Face Pull", "shrug": "Shrug",  
}


# --------------------------------------------------------------- lifecycle

def init():  
    """Create or open the workbook, ensuring all sheets exist."""  
    global _wb, _mtime  
    with _lock:  
        config.BACKUP_DIR.mkdir(exist_ok=True)  
        if config.WORKBOOK_PATH.exists():  
            _wb = openpyxl.load_workbook(config.WORKBOOK_PATH)  
            changed = False  
            for name, headers in _SHEETS:  
                if name not in _wb.sheetnames:  
                    _wb.create_sheet(name).append(headers)  
                    changed = True  
            if changed:  
                _save()  
        else:  
            _wb = openpyxl.Workbook()  
            _wb.remove(_wb.active)  
            for name, headers in _SHEETS:  
                ws = _wb.create_sheet(name)  
                ws.append(headers)  
                ws.freeze_panes = "A2"  
            _seed_program()  
            _save()  
        _mtime = config.WORKBOOK_PATH.stat().st_mtime  
        _backup()  
        _invalidate()


def _seed_program():  
    ws = _wb[config.SHEET_PROGRAM]  
    for row in [  
        ["Monday", 1, "Bench Press", 3, 8, "Top set then back-offs", ""],  
        ["Monday", 2, "Barbell Row", 4, 8, "", ""],  
        ["Monday", 3, "Dip", 3, 10, "Bodyweight", ""],  
        ["Wednesday", 1, "Back Squat", 3, 5, "", ""],  
        ["Wednesday", 2, "Romanian Deadlift", 3, 8, "", ""],  
        ["Friday", 1, "Deadlift", 1, 5, "Work up to a heavy single set", ""],  
        ["Friday", 2, "Overhead Press", 3, 8, "", ""],  
        ["Friday", 3, "Pull Up", 4, 6, "Add weight if easy", ""],  
    ]:  
        ws.append(row)


def _backup():  
    if config.WORKBOOK_PATH.exists():  
        dest = config.BACKUP_DIR / f"workout_{datetime.now():%Y%m%d}.xlsx"  
        if not dest.exists():  
            shutil.copy2(config.WORKBOOK_PATH, dest)


def _save():  
    """Atomic write: temp file then rename. Survives a yanked power cord."""  
    global _mtime  
    tmp = config.WORKBOOK_PATH.with_suffix(".xlsx.tmp")  
    _wb.save(tmp)  
    os.replace(tmp, config.WORKBOOK_PATH)  
    _mtime = config.WORKBOOK_PATH.stat().st_mtime  
    _invalidate()


def _invalidate():  
    global _exercise_cache  
    _exercise_cache = None


def _refresh():  
    """Reload only if the file changed on disk (you edited it in Excel)."""  
    global _wb, _mtime  
    try:  
        m = config.WORKBOOK_PATH.stat().st_mtime  
    except FileNotFoundError:  
        return  
    if _mtime is None or m > _mtime + 0.01:  
        _wb = openpyxl.load_workbook(config.WORKBOOK_PATH)  
        _mtime = m  
        _invalidate()


def _read(sheet):  
    """Rows as dicts, header excluded, blank rows skipped. No disk check."""  
    ws = _wb[sheet]  
    data = list(ws.iter_rows(values_only=True))  
    if not data:  
        return []  
    headers, body = data[0], data[1:]  
    return [dict(zip(headers, r)) for r in body  
            if not all(c is None or str(c).strip() == "" for c in r)]


def rows(sheet):  
    """Public read: checks disk for external edits first."""  
    with _lock:  
        _refresh()  
        return _read(sheet)


# ---------------------------------------------------------- exercise names

def _norm(s):  
    s = re.sub(r"[^a-z0-9 ]+", " ", (s or "").lower())  
    return re.sub(r"\s+", " ", s).strip()


def _alias_map():  
    global _exercise_cache  
    if _exercise_cache is None:  
        m = dict(BUILTIN_ALIASES)  
        for r in _read(config.SHEET_ALIASES):  
            if r.get("Alias") and r.get("Canonical Exercise"):  
                m[_norm(r["Alias"])] = str(r["Canonical Exercise"]).strip()  
        for sheet in (config.SHEET_LOG, config.SHEET_PRS, config.SHEET_PROGRAM):  
            for r in _read(sheet):  
                name = r.get("Exercise")  
                if name:  
                    m.setdefault(_norm(name), str(name).strip())  
        _exercise_cache = m  
    return _exercise_cache


def canonicalize(name):  
    """Map a loose exercise name to its canonical form."""  
    key = _norm(name)  
    if not key:  
        return None  
    with _lock:  
        _refresh()  
        m = _alias_map()  
        if key in m:  
            return m[key]  
        if key.endswith("s") and key[:-1] in m:  
            return m[key[:-1]]  
    return " ".join(w.capitalize() for w in key.split())


def known_exercises():  
    with _lock:  
        _refresh()  
        return sorted(set(_alias_map().values()))


# --------------------------------------------------------------- write ops

def log_set(exercise, sets, reps, weight, unit=None,  
            weight_type="absolute", notes=""):  
    now = datetime.now()  
    unit = unit or config.DEFAULT_UNIT  
    with _lock:  
        _refresh()  
        _wb[config.SHEET_LOG].append([  
            f"{now:%Y-%m-%d}", f"{now:%H:%M}", exercise,  
            int(sets), int(reps), float(weight), unit, weight_type, notes,  
        ])  
        is_pr, previous = _update_pr(exercise, sets, reps, weight, unit, now)  
        _save()  
    return {"exercise": exercise, "sets": int(sets), "reps": int(reps),  
            "weight": float(weight), "unit": unit,  
            "is_pr": is_pr, "previous": previous}


def delete_last():  
    """Undo the newest log row and recompute that PR from remaining history."""  
    with _lock:  
        _refresh()  
        ws = _wb[config.SHEET_LOG]  
        if ws.max_row < 2:  
            return None  
        idx = ws.max_row  
        vals = [c.value for c in ws[idx]]  
        removed = {"exercise": vals[2], "sets": vals[3],  
                   "reps": vals[4], "weight": vals[5]}  
        ws.delete_rows(idx)  
        _recompute_pr(removed["exercise"], removed["sets"], removed["reps"])  
        _save()  
        return removed


# ------------------------------------------------------------------ reads

def recent(limit=12):  
    return list(reversed(rows(config.SHEET_LOG)[-limit:]))


def history_for(exercise, sets=None, reps=None, limit=15):  
    out = []  
    for r in rows(config.SHEET_LOG):  
        if _norm(r.get("Exercise")) != _norm(exercise):  
            continue  
        if sets is not None and r.get("Sets") != sets:  
            continue  
        if reps is not None and r.get("Reps") != reps:  
            continue  
        out.append(r)  
    return list(reversed(out[-limit:]))


def today():  
    stamp = f"{datetime.now():%Y-%m-%d}"  
    return [r for r in rows(config.SHEET_LOG)  
            if str(r.get("Date", "")).startswith(stamp)]


# -------------------------------------------------------------------- PRs

def _find_pr_row(exercise, sets, reps):  
    ws = _wb[config.SHEET_PRS]  
    target = (_norm(exercise), int(sets), int(reps))  
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):  
        if not any(row):  
            continue  
        if (_norm(row[0]), row[1], row[2]) == target:  
            return i  
    return None


def _update_pr(exercise, sets, reps, weight, unit, when):  
    """Returns (is_pr, previous_weight). Caller holds the lock and saves."""  
    ws = _wb[config.SHEET_PRS]  
    sets, reps, weight = int(sets), int(reps), float(weight)  
    i = _find_pr_row(exercise, sets, reps)  
    if i is None:  
        ws.append([exercise, sets, reps, weight, unit, f"{when:%Y-%m-%d}", None])  
        return True, None  
    prev = ws.cell(i, 4).value  
    prev = float(prev) if prev is not None else None  
    if prev is None or weight > prev:  
        ws.cell(i, 4, weight)  
        ws.cell(i, 5, unit)  
        ws.cell(i, 6, f"{when:%Y-%m-%d}")  
        ws.cell(i, 7, prev)  
        return True, prev  
    return False, prev


def _recompute_pr(exercise, sets, reps):  
    """Rebuild one PR from the log. Safer than reversing an update."""  
    matches = [r for r in _read(config.SHEET_LOG)  
               if _norm(r.get("Exercise")) == _norm(exercise)  
               and r.get("Sets") == sets and r.get("Reps") == reps]  
    ws = _wb[config.SHEET_PRS]  
    i = _find_pr_row(exercise, sets, reps)  
    if not matches:  
        if i:  
            ws.delete_rows(i)  
        return  
    best = max(matches, key=lambda r: float(r.get("Weight") or 0))  
    if i is None:  
        ws.append([exercise, sets, reps, float(best["Weight"]),  
                   best.get("Unit") or config.DEFAULT_UNIT, best.get("Date"), None])  
    else:  
        ws.cell(i, 4, float(best["Weight"]))  
        ws.cell(i, 6, best.get("Date"))  
        ws.cell(i, 7, None)


def get_pr(exercise, sets, reps):  
    target = (_norm(exercise), int(sets), int(reps))  
    for r in rows(config.SHEET_PRS):  
        if (_norm(r.get("Exercise")), r.get("Sets"), r.get("Reps")) == target:  
            return r  
    return None


def all_prs(exercise=None):  
    data = rows(config.SHEET_PRS)  
    if exercise:  
        data = [r for r in data if _norm(r.get("Exercise")) == _norm(exercise)]  
    return sorted(data, key=lambda r: (str(r.get("Exercise")),  
                                       r.get("Sets") or 0, r.get("Reps") or 0))


# ---------------------------------------------------------------- program

_DAY_ORDER = {d: i for i, d in enumerate(  
    ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"])}


def program_for_day(day=None):  
    day = (day or f"{datetime.now():%A}").strip().lower()  
    out = [r for r in rows(config.SHEET_PROGRAM)  
           if str(r.get("Day", "")).strip().lower() == day]  
    return sorted(out, key=lambda r: r.get("Order") or 0)


def program_for_exercise(exercise):  
    return [r for r in rows(config.SHEET_PROGRAM)  
            if _norm(r.get("Exercise")) == _norm(exercise)]


def program_all():  
    return sorted(rows(config.SHEET_PROGRAM),  
                  key=lambda r: (_DAY_ORDER.get(str(r.get("Day", "")).lower(), 99),  
                                 r.get("Order") or 0))  